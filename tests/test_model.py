import csv
import io
import json
from pathlib import Path
import subprocess
import sys

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
import pytest

from worldmodel.api import app
from worldmodel.contracts import Baseline, Country, Definitions, Scenario
from worldmodel.engine import simulate
from worldmodel.storage import canonical, export_run, load_definitions, load_run, record_run, sha


@pytest.fixture
def scenario():
    return load_definitions().default_scenario


@pytest.fixture
def country():
    return Country(id="TEST", name="Accounting example", kind="country", members=[], membership_date=None,
                   baseline=Baseline(year=2024, government_scope="General government",
                                     compulsory_revenue_pct_gdp=25, other_revenue_pct_gdp=5,
                                     expenditure_pct_gdp=40, source_ids=[], notes=[]), research_gaps=[])


def changed(scenario, **values):
    return Scenario.model_validate({**scenario.model_dump(), **values})


def test_complete_transition_accounting(scenario, country):
    points = simulate(scenario, [country])[0].points
    assert len(points) == 121
    assert points[0].financing == 10
    assert points[0].funding_gap == 0
    assert points[42].funding_gap == 5
    assert points[60].financing == 0
    assert points[-1].funding_gap == 10
    assert points[-1].sustainable_spending == 30
    for p in points:
        assert p.funded_spending + p.funding_gap == pytest.approx(p.planned_spending)
        assert p.compulsory_revenue / p.gdp_index * 100 <= 25 + 1e-12
        assert p.refinancing_requirement is None


def test_tax_collection_and_declining_output(scenario, country):
    result = simulate(changed(scenario, real_gdp_growth=-0.05, collection_effectiveness=0.6), [country])[0]
    final = result.points[-1]
    assert final.compulsory_revenue == pytest.approx(25 * 0.95 ** 10 * 0.6)
    assert final.compulsory_revenue / final.gdp_index * 100 == pytest.approx(15)
    assert final.planned_spending == 40
    assert final.funding_gap > 10


def test_shared_inflation_and_currency_conversion(scenario, country):
    final = simulate(scenario, [country])[0].points[-1]
    assert final.price_index == pytest.approx(100 * 1.05 ** 10)
    assert final.bitcoin_real_index == pytest.approx(100 * 1.1 ** 10)
    assert final.bitcoin_dollar_index == pytest.approx(100 * 1.05 ** 10 * 1.1 ** 10)
    assert final.stock_dollar_index / final.price_index * 100 == pytest.approx(20)
    assert final.stock_bitcoin_index * final.bitcoin_real_index / 100 == pytest.approx(20)
    alternative = simulate(changed(scenario, annual_inflation=0.25, bitcoin_real_growth=0.4), [country])[0].points[-1]
    assert alternative.funding_gap == final.funding_gap


def test_prices_capacity_and_output_remain_separate(scenario, country):
    final = simulate(scenario, [country])[0].points[-1]
    assert final.stock_real_index == pytest.approx(20)
    assert final.property_real_index == pytest.approx(20)
    assert final.productive_capacity_index == 100
    assert final.gdp_index == 100
    alternative = simulate(changed(scenario, asset_decline=0.99), [country])[0].points[-1]
    assert alternative.stock_real_index == pytest.approx(1)
    assert alternative.gdp_index == final.gdp_index
    assert alternative.funding_gap == final.funding_gap


def test_transition_can_extend_beyond_horizon(scenario, country):
    late = simulate(changed(scenario, transition_onset_month=108, transition_duration_months=36), [country])[0]
    assert late.points[-1].transition == pytest.approx(1 / 3)
    assert late.points[-1].stock_real_index > 20
    absent = simulate(changed(scenario, transition_onset_month=120), [country])[0]
    assert absent.points[-1].funding_gap == 0


def test_surplus_is_not_a_negative_gap(scenario, country):
    country = country.model_copy(update={"baseline": country.baseline.model_copy(update={"expenditure_pct_gdp": 20})})
    final = simulate(scenario, [country])[0].points[-1]
    assert final.sustainable_spending == 30
    assert final.funded_spending == 20
    assert final.financing == final.funding_gap == 0


def test_unknown_fiscal_inputs_stay_unknown(scenario, country):
    unknown = country.model_copy(update={"baseline": None, "research_gaps": ["Missing coverage"]})
    final = simulate(scenario, [unknown])[0].points[-1]
    for key in ("compulsory_revenue", "other_revenue", "planned_spending", "funding_gap", "funding_gap_pct_gdp"):
        assert getattr(final, key) is None
    assert final.stock_real_index == pytest.approx(20)
    assert final.refinancing_requirement is None


def test_deterministic_calculation(scenario, country):
    assert simulate(scenario, [country]) == simulate(scenario, [country])


@pytest.mark.parametrize("values", [
    {"asset_decline": 0.79}, {"collection_effectiveness": 1.1},
    {"transition_duration_months": 0}, {"annual_inflation": float("nan")},
    {"name": ""}, {"exploratory": False}, {"unknown_parameter": 3},
    {"forecast_origin": "2026-10-09"}, {"forecast_origin": "9990-01-01"},
])
def test_invalid_assumptions_rejected(scenario, values):
    with pytest.raises(ValidationError):
        changed(scenario, **values)


def test_saved_scenario_does_not_gain_hidden_defaults(scenario):
    data = scenario.model_dump()
    del data["annual_inflation"]
    with pytest.raises(ValidationError):
        Scenario.model_validate(data)
    assert changed(scenario, name="\u2067العربية 中文 \U0001f30d " + "x" * 100).name


def test_reviewed_coverage_and_bloc_membership():
    definitions = load_definitions()
    eu = next(c for c in definitions.countries if c.id == "EU")
    assert len(eu.members) == len(set(eu.members)) == 27
    assert {"DEU", "FRA"} <= set(eu.members)
    assert "EU" not in eu.members
    assert len(definitions.countries) == 15
    assert all(len(source.sha256) == 64 for source in definitions.sources)
    assert all(country.baseline or country.research_gaps for country in definitions.countries)
    assert {o for m in definitions.mechanisms for o in m.outputs} == {v.id for v in definitions.variables}


def test_causal_registry_rejects_missing_dependencies_and_outputs():
    data = load_definitions().model_dump(mode="json")
    data["mechanisms"][0]["inputs"].append("missing-mechanism")
    with pytest.raises(ValidationError, match="Unknown computational dependency"):
        Definitions.model_validate(data)
    data = load_definitions().model_dump(mode="json")
    data["mechanisms"][0]["outputs"].append("missing-variable")
    with pytest.raises(ValidationError, match="Output must reference"):
        Definitions.model_validate(data)


def test_immutable_runs_and_matching_exports(tmp_path, monkeypatch, scenario):
    monkeypatch.setenv("FUTURE_STATE_DIR", str(tmp_path))
    first = record_run(scenario)
    original = (tmp_path / "runs" / first.run_id / "run.json").read_bytes()
    second = record_run(changed(scenario, real_gdp_growth=-0.02))
    assert first.run_id != second.run_id
    assert (tmp_path / "runs" / first.run_id / "run.json").read_bytes() == original
    assert load_run(first.run_id) == first
    assert first.provenance.scenario_sha256 == sha(canonical(scenario.model_dump(mode="json")))
    workbook = load_workbook(io.BytesIO(export_run(first, "xlsx")), data_only=False)
    assert set(workbook.sheetnames) == {"Trajectories", "Assumptions", "Research gaps", "Sources", "Variables"}
    rows = list(workbook["Trajectories"].values)
    header = rows[0]
    last_us = next(row for row in rows[1:] if row[1] == "USA" and row[header.index("month")] == 120)
    assert last_us[header.index("funding_gap")] == pytest.approx(first.countries[0].points[-1].funding_gap)
    paraguay = next(row for row in rows[1:] if row[1] == "PRY")
    assert paraguay[header.index("funding_gap")] is None
    csv_rows = list(csv.reader(io.StringIO(export_run(first, "csv").decode())))
    header_index = next(i for i, row in enumerate(csv_rows) if row and row[0] == "run_id")
    csv_header = csv_rows[header_index]
    csv_us = next(row for row in csv_rows[header_index + 1:] if len(row) == len(csv_header) and row[1] == "USA" and row[csv_header.index("month")] == "120")
    assert float(csv_us[csv_header.index("funding_gap")]) == first.countries[0].points[-1].funding_gap
    assert next(row for row in csv_rows if row and row[0] == "# Variables / run_id")
    combined = load_workbook(io.BytesIO(export_run(first, "xlsx", second)), data_only=True)
    combined_rows = list(combined["Trajectories"].values)
    assert len(combined_rows) == 2 * (len(rows) - 1) + 1
    assert {row[0] for row in combined_rows[1:]} == {first.run_id, second.run_id}
    assert dict(combined["Assumptions"].values)["comparison_run_id"] == second.run_id


def test_api_validation_and_exports(tmp_path, monkeypatch, scenario):
    monkeypatch.setenv("FUTURE_STATE_DIR", str(tmp_path))
    client = TestClient(app)
    assert client.get("/api/definitions").status_code == 200
    assert client.post("/api/runs", json={}).status_code == 422
    assert client.post("/api/runs", json={**scenario.model_dump(mode="json"), "asset_decline": 0.5}).status_code == 422
    response = client.post("/api/runs", json=scenario.model_dump(mode="json"))
    assert response.status_code == 200
    run_id = response.json()["run_id"]
    export = client.get(f"/api/runs/{run_id}/export?format=xlsx")
    assert export.status_code == 200
    assert len(export.content) > 10000
    compared = client.post("/api/runs", json=changed(scenario, collection_effectiveness=0.5).model_dump(mode="json")).json()
    combined = client.get(f"/api/runs/{run_id}/export?format=xlsx&comparison_run_id={compared['run_id']}")
    assert combined.status_code == 200
    workbook = load_workbook(io.BytesIO(combined.content), data_only=True)
    assert workbook["Trajectories"].max_row == 15 * 121 * 2 + 1
    assert client.get(f"/api/runs/{run_id}/export?format=csv&comparison_run_id=invalid").status_code == 404
    assert client.get(f"/api/runs/{run_id}/export?format=pdf").status_code == 422
    assert client.get("/api/runs/invalid/export?format=csv").status_code == 404


def test_registry_version_guard_preserves_runs(tmp_path, monkeypatch, scenario):
    monkeypatch.setenv("FUTURE_STATE_DIR", str(tmp_path))
    client = TestClient(app)
    definitions = client.get("/api/definitions")
    etag = definitions.headers["etag"]
    assert etag == f'"{sha(canonical(definitions.json()))}"'
    accepted = client.post("/api/runs", json=scenario.model_dump(mode="json"), headers={"If-Match": etag})
    assert accepted.status_code == 200
    assert accepted.json()["provenance"]["registry_sha256"] == etag.strip('"')
    original_runs = set((tmp_path / "runs").iterdir())
    rejected = client.post("/api/runs", json=scenario.model_dump(mode="json"), headers={"If-Match": '"stale-registry"'})
    assert rejected.status_code == 412
    assert "Reload the page" in rejected.json()["detail"]
    assert set((tmp_path / "runs").iterdir()) == original_runs


def test_failed_cli_export_does_not_create_an_output(tmp_path):
    output = tmp_path / "invalid.csv"
    result = subprocess.run([sys.executable, "-m", "worldmodel", "export", "invalid",
                             "--format", "csv", "--output", str(output)], capture_output=True, text=True)
    assert result.returncode != 0
    assert "Invalid run ID" in result.stderr
    assert not output.exists()
