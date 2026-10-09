"""Collect public source files for manual review. Never modify accepted records.

Run with: uv run python scripts/collect_fiscal_sources.py
The output is an immutable proposal in .cache/fiscal-review/<timestamp>/.
"""

from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import urllib.request

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCES = {
    "world": "https://www.imf.org/-/media/files/topics/fiscal/world-revenue-longitudinal-database/world-imf2026.xlsx",
    "fm_revenue": "https://www.imf.org/external/datamapper/api/v2/GGR_G01_GDP_PT",
    "fm_expenditure": "https://www.imf.org/external/datamapper/api/v2/G_X_G01_GDP_PT",
    "eu": "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/gov_10a_main?lang=EN&geo=EU27_2020&unit=PC_GDP&sector=S13&time=2024",
    "fr": "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/gov_10a_main?lang=EN&geo=FR&unit=PC_GDP&sector=S13&time=2024",
    "de": "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/gov_10a_main?lang=EN&geo=DE&unit=PC_GDP&sector=S13&time=2024",
}
COUNTRIES = ["USA", "CHN", "JPN", "DEU", "GBR", "FRA", "IND", "BRA", "RUS", "SAU", "CHE", "SGP", "ARG", "PRY"]


def main():
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    output = ROOT / ".cache/fiscal-review" / timestamp
    output.mkdir(parents=True, exist_ok=False)

    def fetch(item):
        key, url = item
        body = urllib.request.urlopen(url, timeout=40).read()
        (output / (key + (".xlsx" if key == "world" else ".json"))).write_bytes(body)
        return key, body

    with ThreadPoolExecutor(max_workers=6) as pool:
        files = dict(pool.map(fetch, SOURCES.items()))
    workbook = load_workbook(io.BytesIO(files["world"]), read_only=True, data_only=True)
    rows = workbook["Data"].iter_rows(values_only=True)
    headers = list(next(rows))
    by_country = {c: [] for c in COUNTRIES}
    for values in rows:
        row = dict(zip(headers, values))
        if row.get("ISO3") in by_country and row.get("year", 0) <= 2024:
            by_country[row["ISO3"]].append(row)
    fm_revenue = json.loads(files["fm_revenue"])["values"]["GGR_G01_GDP_PT"]
    fm_exp = json.loads(files["fm_expenditure"])["values"]["G_X_G01_GDP_PT"]
    candidates = {}
    for country, rows in by_country.items():
        candidates[country] = []
        for row in sorted(rows, key=lambda r: r["year"], reverse=True):
            year = int(row["year"])
            if year < 2020 or any(row.get(k) is None for k in ("TotRev", "TaxRev", "SocialCon")):
                continue
            revenue = fm_revenue.get(country, {}).get(str(year))
            expenditure = fm_exp.get(country, {}).get(str(year))
            if revenue is None or expenditure is None:
                continue
            candidates[country].append({
                "year": year, "world_total": row["TotRev"], "tax": row["TaxRev"],
                "social": row["SocialCon"], "fm_total": revenue, "expenditure": expenditure,
                "revenue_difference_pp": revenue - row["TotRev"],
            })
    for key in ("eu", "fr", "de"):
        data = json.loads(files[key])
        items = data["dimension"]["na_item"]["category"]["index"]
        candidates[key] = {k: data["value"].get(str(items[k])) for k in
                           ("D2REC", "D5REC", "D91REC", "D611REC", "D613REC", "TR", "TE")}
    proposal = {
        "retrieved_at": timestamp,
        "sources": {k: {"url": SOURCES[k], "sha256": hashlib.sha256(b).hexdigest()} for k, b in files.items()},
        "candidates": candidates,
    }
    (output / "proposal.json").write_text(json.dumps(proposal, indent=2) + "\n")
    print(output / "proposal.json")


if __name__ == "__main__":
    main()
