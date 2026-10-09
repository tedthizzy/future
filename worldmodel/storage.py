"""Immutable local runs. Calculation and publication share the recorded values."""

import csv
import hashlib
import importlib.metadata
import io
import json
import os
from pathlib import Path
import platform
import subprocess
from datetime import datetime, timezone
from uuid import uuid4

from openpyxl import Workbook

from .contracts import Definitions, Point, Provenance, Run, Scenario
from .engine import simulate

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_STATE = Path.home() / ".local/share/future"


class RegistryChanged(ValueError):
    """The caller's evidence no longer matches the definitions used for execution."""


def canonical(value) -> bytes:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=True).encode()


def sha(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def load_definitions() -> Definitions:
    return Definitions.model_validate_json((ROOT / "registry/definitions.json").read_bytes())


def state_dir() -> Path:
    # This override exists for isolated tests and deliberate private storage choices.
    return Path(os.environ.get("FUTURE_STATE_DIR", str(DEFAULT_STATE)))


def record_run(scenario: Scenario, expected_registry: str | None = None) -> Run:
    definitions = load_definitions()
    registry_hash = sha(canonical(definitions.model_dump(mode="json")))
    if expected_registry is not None and expected_registry != registry_hash:
        raise RegistryChanged("Model definitions changed. Reload the page before calculating.")
    code_files = sorted((ROOT / "worldmodel").glob("*.py"))
    code_hash = sha(b"".join(p.name.encode() + p.read_bytes() for p in code_files))
    revision = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
    dirty = bool(subprocess.check_output(["git", "status", "--porcelain"], cwd=ROOT, text=True).strip())
    run = Run(
        schema_version="1", model_version=definitions.model_version, run_id=uuid4().hex,
        scenario=scenario,
        provenance=Provenance(
            code_revision=revision, code_dirty=dirty, code_sha256=code_hash,
            registry_sha256=registry_hash,
            scenario_sha256=sha(canonical(scenario.model_dump(mode="json"))),
            python_version=platform.python_version(),
            dependencies={k: importlib.metadata.version(k) for k in
                          ("pydantic", "fastapi", "uvicorn", "openpyxl")},
            solver="deterministic monthly accounting; annualized flows",
            created_at=datetime.now(timezone.utc).isoformat(),
        ), countries=simulate(scenario, definitions.countries),
    )
    directory = state_dir() / "runs" / run.run_id
    directory.mkdir(parents=True, exist_ok=False)
    # Write result last; incomplete directories are never presented as complete runs.
    (directory / "resolved-scenario.json").write_bytes(canonical(scenario.model_dump(mode="json")))
    (directory / "definitions.json").write_bytes(canonical(definitions.model_dump(mode="json")))
    (directory / "run.json").write_bytes(canonical(run.model_dump(mode="json")))
    return run


def load_run(run_id: str) -> Run:
    if len(run_id) != 32 or any(c not in "0123456789abcdef" for c in run_id):
        raise ValueError("Invalid run ID")
    return Run.model_validate_json((state_dir() / "runs" / run_id / "run.json").read_bytes())


def export_run(run: Run, format: str, comparison: Run | None = None) -> bytes:
    fields = ["run_id", "country_id", "baseline_year", *Point.model_fields]
    rows = []
    runs = [run] if comparison is None else [run, comparison]
    for recorded in runs:
        for country in recorded.countries:
            for point in country.points:
                rows.append([recorded.run_id, country.country_id, country.baseline.year if country.baseline else None,
                             *point.model_dump(mode="json").values()])
    metadata = [
        ["schema_version", run.schema_version], ["model_version", run.model_version],
        ["scenario_json", canonical(run.scenario.model_dump(mode="json")).decode()],
        ["provenance_json", canonical(run.provenance.model_dump(mode="json")).decode()],
        ["flow_units", "Annualized real units; starting annual GDP = 100"],
        ["index_units", "Starting value = 100; no absolute dollar or Bitcoin price"],
        ["refinancing_requirement", "Unknown; excluded from recurring funding gap"],
    ]
    if comparison is not None:
        metadata.extend([
            ["comparison_run_id", comparison.run_id],
            ["comparison_scenario_json", canonical(comparison.scenario.model_dump(mode="json")).decode()],
            ["comparison_provenance_json", canonical(comparison.provenance.model_dump(mode="json")).decode()],
        ])
    gaps, sources, units = [], [], []
    for recorded in runs:
        definitions = Definitions.model_validate_json(
            (state_dir() / "runs" / recorded.run_id / "definitions.json").read_bytes())
        gaps.extend([[recorded.run_id, c.country_id, "; ".join(c.research_gaps)] for c in recorded.countries])
        sources.extend([[recorded.run_id, s.id, s.title, s.url, s.publication_date, s.retrieval_date, s.sha256, s.notes]
                        for s in definitions.sources])
        units.extend([[recorded.run_id, v.id, v.name, v.unit, v.status, v.mechanism_id] for v in definitions.variables])
    if format == "csv":
        out = io.StringIO(newline="")
        writer = csv.writer(out)
        writer.writerows([["# " + str(k), v] for k, v in metadata])
        writer.writerow(fields)
        writer.writerows(rows)
        writer.writerow([])
        writer.writerow(["# Research gaps / run_id", "Country", "Details"])
        writer.writerows(gaps)
        writer.writerow([])
        writer.writerow(["# Sources / run_id", "ID", "Title", "URL", "Published", "Retrieved", "SHA256", "Notes"])
        writer.writerows(sources)
        writer.writerow([])
        writer.writerow(["# Variables / run_id", "ID", "Name", "Units", "Status", "Mechanism"])
        writer.writerows(units)
        return out.getvalue().encode()
    if format != "xlsx":
        raise ValueError("Export format must be csv or xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Trajectories"
    sheet.append(fields)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    for title, records in [("Assumptions", metadata), ("Research gaps", gaps),
                           ("Sources", sources), ("Variables", units)]:
        sheet = workbook.create_sheet(title)
        for row in records:
            sheet.append(row)
            # Scenario names and source notes are data, never spreadsheet formulas.
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
